"""Copiar o exportar lo que responde el agente.

Las respuestas acaban en un correo, en Google Chat, en una reunión o en una
hoja de seguimiento. De cada turno salen:
- el texto sin formato de Markdown, para pegar en un correo o un chat;
- cada lista que devolvió una herramienta, como tabla para pegar en una hoja
  (columnas separadas por tabuladores) o en un Excel con una hoja por lista.

Las listas salen de los datos de las herramientas, no del texto del modelo.
"""
from __future__ import annotations

import base64
import datetime
import html
import io
import json
import re
from typing import Any, Iterable
from zoneinfo import ZoneInfo

# herramienta -> (clave de la lista en la respuesta, título de la hoja)
_LISTAS = {
    "buscar_propiedades": ("matches", "Villas"),
    "listar_propiedades": ("matches", "Villas"),
    "buscar_por_valoracion": ("matches", "Villas por valoración"),
    "consultar_disponibilidad": ("matches", "Villas disponibles"),
    "buscar_ofertas": ("matches", "Ofertas"),
    "alternativas_villa": ("matches", "Alternativas"),
    "consultar_reservas": ("reservas", "Reservas"),
    "resumen_reservas": ("resumen", "Resumen de reservas"),
    "consultar_precios": ("noches", "Precios por noche"),
    "calendario_villa": ("tramos", "Calendario"),
    "detalle_reserva": ("pagos", "Plan de pagos"),
}
# Claves de la respuesta con el total real, por si la lista está recortada.
_TOTALES = ("total", "total_disponibles", "total_grupos")
# Identificadores internos que no dicen nada a quien lee la hoja.
_OCULTAS = {"villa_id", "reserva_id", "cliente_id", "huesped_id", "ficha_id", "total_resultados"}

# Columnas que van al Excel y a la copia, en este orden. Lo que no está aquí
# (identificadores, moneda, banderas internas) no se exporta.
_COLUMNAS = [
    ("localizador", "Localizador"), ("cliente_nombre", "Titular"), ("titular", "Titular"),
    ("nombre", "Villa"), ("villa_nombre", "Villa"), ("villa", "Villa"),
    ("pueblo_cercano", "Pueblo"), ("zona", "Zona"), ("direccion", "Dirección"),
    ("dimension", "Grupo"),
    ("fecha", "Noche"), ("desde", "Desde"), ("hasta", "Hasta"),
    ("fecha_entrada", "Entrada"), ("fecha_salida", "Salida"), ("noches", "Noches"),
    ("adultos", "Adultos"), ("ninos", "Niños"), ("num_mascotas", "Mascotas"),
    ("capacidad_pax", "Personas"), ("numero_habitaciones", "Habitaciones"),
    ("numero_banos", "Baños"), ("camas_totales", "Camas"), ("m2_habitables", "m² habitables"),
    ("tiene_piscina_privada", "Piscina privada"), ("admite_animales", "Admite mascotas"),
    ("tiene_vista_mar", "Vistas al mar"),
    ("distancia_mar_m", "Distancia al mar (m)"),
    ("distancia_playa_arena_m", "Distancia a playa de arena (m)"),
    ("rating_medio", "Valoración media"),
    ("tipo_ocupacion", "Estado"), ("subtipo_reserva", "Tipo"), ("estado_reserva", "Estado"),
    ("estado_limpieza", "Limpieza"),
    ("precio_total", "Precio total (€)"), ("precio_medio_noche", "Precio medio por noche (€)"),
    ("precio_venta", "Venta (€)"), ("precio_compra", "Compra (€)"),
    ("margen", "Margen (€)"), ("margen_total", "Margen (€)"), ("margen_pct", "Margen (%)"),
    ("larga_estancia", "Larga estancia"),
    ("total_reservas", "Reservas"), ("importe_total", "Importe total (€)"),
    ("importe_medio", "Importe medio (€)"), ("importe_por_noche", "Importe por noche (€)"),
    ("noches_totales", "Noches totales"), ("noches_medias", "Noches medias"),
    ("fecha_vencimiento", "Vencimiento"), ("importe_previsto", "Previsto (€)"),
    ("importe_pagado", "Pagado (€)"), ("importe_pendiente", "Pendiente (€)"),
    ("metodo_pago_nombre", "Forma de pago"), ("fecha_pago", "Pagado el"),
    ("fecha_confirmacion", "Confirmada el"), ("fecha_pedido", "Hecha el"),
    ("fecha_anulacion", "Anulada el"),
]
_ETIQUETAS = dict(_COLUMNAS)
_ORDEN = {clave: i for i, (clave, _) in enumerate(_COLUMNAS)}
_FECHAS = {"fecha", "desde", "hasta", "fecha_entrada", "fecha_salida", "fecha_vencimiento",
           "fecha_pago", "fecha_confirmacion", "fecha_pedido", "fecha_anulacion"}
_EUROS = {clave for clave, etiqueta in _COLUMNAS if "(€)" in etiqueta}
_DECIMALES = {"margen_pct", "rating_medio", "noches_medias"}

_TINTA = "1E2B3C"          # azul marino de la marca
_TINTA_SUAVE = "4A5768"
_FILA_ALTERNA = "F5F2E9"   # crema de la marca


def _celda(valor: Any) -> Any:
    if isinstance(valor, (list, tuple, set)):
        return ", ".join(str(v) for v in valor)
    if isinstance(valor, dict):
        return ", ".join(f"{k}: {v}" for k, v in valor.items())
    return valor


def tablas(llamadas: Iterable[tuple[str, dict, Any]]) -> list[dict]:
    """Cada lista con datos que devolvieron las herramientas del turno."""
    resultado = []
    for nombre, _, respuesta in llamadas:
        if nombre not in _LISTAS or not isinstance(respuesta, dict) or respuesta.get("error"):
            continue
        clave, titulo = _LISTAS[nombre]
        filas = [
            {k: _celda(v) for k, v in fila.items() if k not in _OCULTAS}
            for fila in respuesta.get(clave) or [] if isinstance(fila, dict)
        ]
        if not filas:
            continue
        total = next((respuesta[k] for k in _TOTALES if isinstance(respuesta.get(k), int)), None)
        resultado.append({"titulo": titulo, "filas": filas, "total": total})
    return resultado


def _columnas(tabla: dict) -> list[str]:
    """Las columnas exportables de una tabla, en el orden de _COLUMNAS, sin
    repetir etiqueta (una hoja no necesita "Villa" dos veces)."""
    # Una columna vacía en todas las filas (p. ej. "Anulada el" en reservas
    # vigentes) solo estorba.
    presentes = {c for fila in tabla["filas"] for c, v in fila.items()
                 if c in _ETIQUETAS and v not in (None, "")}
    columnas, vistas = [], set()
    for clave in sorted(presentes, key=_ORDEN.get):
        if _ETIQUETAS[clave] not in vistas:
            vistas.add(_ETIQUETAS[clave])
            columnas.append(clave)
    return columnas


def _valor(clave: str, valor: Any) -> Any:
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if clave in _FECHAS and valor:
        try:
            return datetime.datetime.fromisoformat(str(valor)[:10])
        except ValueError:
            return valor
    return valor


def a_tsv(tabla: dict) -> str:
    """Columnas separadas por tabuladores: al pegar en Google Sheets o Excel
    cada dato cae en su celda."""
    columnas = _columnas(tabla)
    lineas = ["\t".join(_ETIQUETAS[c] for c in columnas)]
    for fila in tabla["filas"]:
        celdas = []
        for c in columnas:
            v = _valor(c, fila.get(c))
            celdas.append(v.strftime("%d/%m/%Y") if isinstance(v, datetime.datetime)
                          else "" if v is None else str(v))
        lineas.append("\t".join(celdas))
    return "\n".join(lineas)


def a_excel(lista_tablas: list[dict], generado: datetime.datetime | None = None) -> bytes:
    """Un .xlsx ordenado, con una hoja por lista: título y resumen arriba,
    cabecera destacada, fija y con filtros, y cada dato con su formato."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    generado = generado or datetime.datetime.now(ZoneInfo("Europe/Madrid"))
    libro = Workbook()
    libro.remove(libro.active)
    usados: dict[str, int] = {}
    linea = Side(style="thin", color="E2E0D6")
    for tabla in lista_tablas:
        base = tabla["titulo"][:26]
        usados[base] = usados.get(base, 0) + 1
        hoja = libro.create_sheet(base if usados[base] == 1 else f"{base} ({usados[base]})")
        columnas = _columnas(tabla)
        n = len(tabla["filas"])

        hoja["A1"] = tabla["titulo"]
        hoja["A1"].font = Font(bold=True, size=14, color=_TINTA)
        resumen = f"{n} {'fila' if n == 1 else 'filas'}"
        if tabla.get("total") and tabla["total"] > n:
            resumen = f"{n} de {tabla['total']} que cumplen la búsqueda"
        hoja["A2"] = f"{resumen} · generado el {generado:%d/%m/%Y %H:%M} · Asistente Abahana Villas"
        hoja["A2"].font = Font(italic=True, size=9, color=_TINTA_SUAVE)

        for j, clave in enumerate(columnas, start=1):
            celda = hoja.cell(row=4, column=j, value=_ETIQUETAS[clave])
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="FF" + _TINTA)
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.row_dimensions[4].height = 30

        for i, fila in enumerate(tabla["filas"], start=5):
            for j, clave in enumerate(columnas, start=1):
                celda = hoja.cell(row=i, column=j, value=_valor(clave, fila.get(clave)))
                celda.border = Border(bottom=linea)
                if i % 2 == 0:
                    celda.fill = PatternFill("solid", fgColor="FF" + _FILA_ALTERNA)
                if isinstance(celda.value, datetime.datetime):
                    celda.number_format = "dd/mm/yyyy"
                elif clave in _EUROS:
                    celda.number_format = '#,##0.00 "€"'
                elif clave in _DECIMALES:
                    celda.number_format = "0.0"

        if columnas:
            ultima = get_column_letter(len(columnas))
            hoja.freeze_panes = "A5"
            hoja.auto_filter.ref = f"A4:{ultima}{4 + n}"
            for j, clave in enumerate(columnas, start=1):
                largos = [len(_ETIQUETAS[clave])] + [
                    10 if clave in _FECHAS else len(str(fila.get(clave) or ""))
                    for fila in tabla["filas"][:200]]
                hoja.column_dimensions[get_column_letter(j)].width = min(max(largos) + 3, 45)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def botones_html(texto_markdown: str, excel: bytes | None, nombre_excel: str) -> str:
    """"Copiar" (al portapapeles, en un clic) y "Excel" (descarga), alineados
    abajo a la derecha con los colores de la marca.

    Va en un componente HTML porque Streamlit no tiene un botón que copie al
    portapapeles; su iframe permite clipboard-write y descargas. El texto va
    como JSON y sin "</" para que no pueda cerrar el <script>.
    """
    texto = json.dumps(texto_plano(texto_markdown)).replace("</", "<\\/")
    boton_excel = ""
    if excel:
        datos = base64.b64encode(excel).decode()
        boton_excel = (
            f'<a class="b" download="{html.escape(nombre_excel)}" '
            f'href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{datos}">'
            '<span class="i">⤓</span>Excel</a>')
    return f"""<!doctype html><html><head><meta charset="utf-8">
<link href="https://fonts.googleapis.com/css2?family=Montserrat:wght@500;600&display=swap" rel="stylesheet">
<style>
  body {{ margin: 0; background: transparent; font-family: 'Montserrat', sans-serif; }}
  html, body {{ height: 100%; }}
  .fila {{ display: flex; justify-content: flex-end; align-items: center; gap: 8px;
          height: 100%; padding-right: 2px; box-sizing: border-box; }}
  .b {{ display: inline-flex; align-items: center; gap: 6px; height: 32px; padding: 0 12px;
        border: 1px solid #E2E0D6; border-radius: 8px; background: #FFFFFF; color: #1E2B3C;
        font: 500 13px 'Montserrat', sans-serif; text-decoration: none; cursor: pointer; }}
  .b:hover {{ border-color: #7089C4; color: #3E5A9A; background: #E9EEF8; }}
  .b:focus-visible {{ outline: 2px solid #3E5A9A; outline-offset: 2px; }}
  .i {{ font-size: 14px; line-height: 1; }}
</style></head><body><div class="fila">
  <button class="b" id="copiar" type="button" title="Copiar la respuesta sin formato, para un correo o un chat">
    <span class="i">⧉</span><span id="etiqueta">Copiar</span></button>
  {boton_excel}
</div><script>
  const texto = {texto};
  const boton = document.getElementById('copiar');
  const etiqueta = document.getElementById('etiqueta');
  function avisar(msg) {{ etiqueta.textContent = msg; setTimeout(() => etiqueta.textContent = 'Copiar', 1800); }}
  boton.addEventListener('click', async () => {{
    try {{ await navigator.clipboard.writeText(texto); avisar('Copiado'); }}
    catch (e) {{
      const area = document.createElement('textarea'); area.value = texto;
      document.body.appendChild(area); area.select();
      const ok = document.execCommand('copy'); area.remove();
      avisar(ok ? 'Copiado' : 'No se pudo copiar');
    }}
  }});
</script></body></html>"""


def texto_plano(markdown: str) -> str:
    """La respuesta sin la sintaxis de Markdown, para un correo o un chat."""
    lineas = []
    en_bloque_codigo = False
    for linea in (markdown or "").splitlines():
        if linea.strip().startswith("```"):
            en_bloque_codigo = not en_bloque_codigo
            continue
        if not en_bloque_codigo:
            if re.fullmatch(r"\s*([-*_])\1{2,}\s*", linea):
                continue
            linea = re.sub(r"^\s{0,3}#{1,6}\s*", "", linea)
            linea = re.sub(r"^(\s*)[*\-+]\s+", r"\1• ", linea)
            linea = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), linea)
            linea = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])", r"\1", linea)
            linea = re.sub(r"`([^`]*)`", r"\1", linea)
            linea = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", linea)
        lineas.append(linea.rstrip())
    texto = "\n".join(lineas)
    return re.sub(r"\n{3,}", "\n\n", texto).strip()
